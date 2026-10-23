import json
import uuid
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook

DATA_DIR = Path(__file__).parent / "data"
EXCEL_PATH = DATA_DIR / "extractions.xlsx"

COLUMNS = [
    "extraction_id",
    "timestamp",
    "equipment_type",
    "image_filename",
    "fields_json",
    "review_flags_json",
    "status",
    "user_email",
]


def _ensure_workbook():
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    if not EXCEL_PATH.exists():
        wb = Workbook()
        ws = wb.active
        ws.title = "extractions"
        ws.append(COLUMNS)
        wb.save(EXCEL_PATH)


def save_extraction(
    equipment_type: str,
    image_filename: str,
    fields: dict,
    review_flags: dict,
    user_email: str = "",
) -> str:
    _ensure_workbook()
    wb = load_workbook(EXCEL_PATH)
    ws = wb["extractions"]

    extraction_id = str(uuid.uuid4())
    timestamp = datetime.now().isoformat()

    ws.append([
        extraction_id,
        timestamp,
        equipment_type,
        image_filename,
        json.dumps(fields),
        json.dumps(review_flags),
        "pending",
        user_email,
    ])

    wb.save(EXCEL_PATH)
    return extraction_id


def get_all_extractions() -> list[dict]:
    _ensure_workbook()
    wb = load_workbook(EXCEL_PATH)
    ws = wb["extractions"]

    rows = list(ws.iter_rows(min_row=2, values_only=True))
    results = []
    for row in rows:
        if row[0] is None:
            continue
        results.append({
            "extraction_id": row[0],
            "timestamp": row[1],
            "equipment_type": row[2],
            "image_filename": row[3],
            "fields": json.loads(row[4]) if row[4] else {},
            "review_flags": json.loads(row[5]) if row[5] else {},
            "status": row[6],
            "user_email": row[7] if len(row) > 7 and row[7] else "",
        })
    return results


def get_extraction_by_id(extraction_id: str) -> dict | None:
    all_extractions = get_all_extractions()
    for ext in all_extractions:
        if ext["extraction_id"] == extraction_id:
            return ext
    return None


def update_extraction_fields(extraction_id: str, corrected_fields: dict) -> bool:
    _ensure_workbook()
    wb = load_workbook(EXCEL_PATH)
    ws = wb["extractions"]

    for row in ws.iter_rows(min_row=2):
        if row[0].value == extraction_id:
            existing_fields = json.loads(row[4].value) if row[4].value else {}
            existing_fields.update(corrected_fields)
            row[4].value = json.dumps(existing_fields)
            row[6].value = "reviewed"
            wb.save(EXCEL_PATH)
            return True
    return False

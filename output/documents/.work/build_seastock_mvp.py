from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path('/home/suchithgoud/Downloads/Label_Reader')
OUT = ROOT / 'output/documents/Seastock_MVP_Requirements_and_Demo_Flow.docx'
doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(.65)
section.bottom_margin = Inches(.65)
section.left_margin = Inches(.75)
section.right_margin = Inches(.75)
section.footer_distance = Inches(.25)

for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3', 'List Number', 'List Bullet', 'List Paragraph']:
    style = doc.styles[name]
    style.font.name = 'Arial'
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.widow_control = True
for style in doc.styles:
    for border in list(style.element.xpath('.//w:pBdr')):
        border.getparent().remove(border)
normal = doc.styles['Normal']
normal.font.size = Pt(11)
normal.paragraph_format.line_spacing = 1.10
normal.paragraph_format.space_after = Pt(7)
doc.styles['Title'].font.size = Pt(25)
doc.styles['Title'].font.bold = True
doc.styles['Title'].paragraph_format.space_after = Pt(10)
doc.styles['Subtitle'].font.size = Pt(12)
doc.styles['Subtitle'].paragraph_format.space_after = Pt(8)
for name, size in [('Heading 1', 18), ('Heading 2', 12.5), ('Heading 3', 11.5)]:
    doc.styles[name].font.size = Pt(size)
    doc.styles[name].font.bold = True
    doc.styles[name].paragraph_format.space_before = Pt(10)
    doc.styles[name].paragraph_format.space_after = Pt(7)
    doc.styles[name].paragraph_format.keep_with_next = True

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
run = footer.add_run('Pixorex Seastock MVP   ')
run.font.size = Pt(9)
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)
doc.core_properties.title = 'Seastock MVP requirements and demo flow'
doc.core_properties.subject = 'Delivery note scanning and can receipt and issue inventory demo'
doc.core_properties.author = 'Pixorex'
doc.core_properties.keywords = 'Seastock, Pixorex, article master, physical stock, delivery note, paint cans'

numbering = doc.part.numbering_part.element
abstract_id = max(int(node.get(qn('w:abstractNumId'))) for node in numbering.findall(qn('w:abstractNum'))) + 1
abstract = OxmlElement('w:abstractNum')
abstract.set(qn('w:abstractNumId'), str(abstract_id))
level = OxmlElement('w:lvl')
level.set(qn('w:ilvl'), '0')
for tag, value in [('start', '1'), ('numFmt', 'decimal'), ('lvlText', '%1.'), ('lvlJc', 'left')]:
    child = OxmlElement('w:' + tag)
    child.set(qn('w:val'), value)
    level.append(child)
indent_props = OxmlElement('w:pPr')
indent = OxmlElement('w:ind')
indent.set(qn('w:left'), '300')
indent.set(qn('w:hanging'), '300')
indent_props.append(indent)
tabs = OxmlElement('w:tabs')
tab = OxmlElement('w:tab')
tab.set(qn('w:val'), 'num')
tab.set(qn('w:pos'), '300')
tabs.append(tab)
indent_props.insert(0, tabs)
level.append(indent_props)
abstract.append(level)
numbering.insert(list(numbering).index(numbering.find(qn('w:num'))), abstract)
current_num_id = None

def p(text, bold=False):
    para = doc.add_paragraph()
    para.add_run(text).bold = bold
    return para

def h(text, level=2):
    return doc.add_heading(text, level)

def step(text):
    global current_num_id
    if current_num_id is None:
        current_num_id = max(int(node.get(qn('w:numId'))) for node in numbering.findall(qn('w:num'))) + 1
        num = OxmlElement('w:num')
        num.set(qn('w:numId'), str(current_num_id))
        reference = OxmlElement('w:abstractNumId')
        reference.set(qn('w:val'), str(abstract_id))
        num.append(reference)
        override = OxmlElement('w:lvlOverride')
        override.set(qn('w:ilvl'), '0')
        start = OxmlElement('w:startOverride')
        start.set(qn('w:val'), '1')
        override.append(start)
        num.append(override)
        numbering.append(num)
    para = doc.add_paragraph(text, style='List Number')
    props = para._p.get_or_add_pPr().get_or_add_numPr()
    props.get_or_add_ilvl().val = 0
    props.get_or_add_numId().val = current_num_id
    para.paragraph_format.space_after = Pt(6)
    return para

def bullet(text):
    para = doc.add_paragraph(text, style='List Bullet')
    para.paragraph_format.space_after = Pt(5)
    return para

def page(title):
    global current_num_id
    current_num_id = None
    doc.add_page_break()
    h(title, 1)

def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for col, width in zip(t.columns, widths):
        col.width = Inches(width)
    for cell, label, width in zip(t.rows[0].cells, headers, widths):
        cell.width = Inches(width)
        cell.text = label
    for row in rows:
        cells = t.add_row().cells
        for cell, text, width in zip(cells, row, widths):
            cell.width = Inches(width)
            cell.text = text
    for index, row in enumerate(t.rows):
        row_properties = row._tr.get_or_add_trPr()
        if index == 0:
            repeat = OxmlElement('w:tblHeader')
            repeat.set(qn('w:val'), 'true')
            row_properties.append(repeat)
        cant_split = OxmlElement('w:cantSplit')
        row_properties.append(cant_split)
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props = cell._tc.get_or_add_tcPr()
            margins = OxmlElement('w:tcMar')
            for edge, value in [('top', 95), ('bottom', 95), ('left', 110), ('right', 110)]:
                node = OxmlElement('w:' + edge)
                node.set(qn('w:w'), str(value))
                node.set(qn('w:type'), 'dxa')
                margins.append(node)
            props.append(margins)
            borders = OxmlElement('w:tcBorders')
            for edge in ['top', 'left', 'bottom', 'right']:
                node = OxmlElement('w:' + edge)
                node.set(qn('w:val'), 'single')
                node.set(qn('w:sz'), '4')
                node.set(qn('w:color'), 'D9D9D9')
                borders.append(node)
            props.append(borders)
            shade = OxmlElement('w:shd')
            shade.set(qn('w:fill'), '0B1F44' if index == 0 else ('F1F5F9' if index % 2 == 0 else 'FFFFFF'))
            props.append(shade)
            for para in cell.paragraphs:
                para.paragraph_format.space_after = Pt(0)
                para.paragraph_format.line_spacing = 1.05
                for run in para.runs:
                    run.font.name = 'Arial'
                    run.font.size = Pt(10)
                    run.font.bold = index == 0
                    run.font.color.rgb = RGBColor.from_string('FFFFFF' if index == 0 else '000000')
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t

def code(lines):
    for line in lines:
        para = doc.add_paragraph()
        para.paragraph_format.space_after = Pt(1)
        para.paragraph_format.line_spacing = 1.0
        run = para.add_run(line)
        run.font.name = 'Liberation Mono'
        run.font.size = Pt(10)
    para.paragraph_format.space_after = Pt(9)

doc.add_paragraph('Seastock MVP requirements and demo flow', style='Title')
doc.add_paragraph('Pixorex   2 October 2026   Can receipt and issue', style='Subtitle')
p('Can scanning is a required MVP capability. The demo scans the Delivery Note, displays expected items, scans the supplied labels on cans to receive stock, and scans the same labels in Issue mode to deduct cans leaving the locker room.')
p('The MVP follows the happy path. If a can label is unclear, the operator enters its article number manually. Special label recovery and automated exception workflows are outside this scope. The assumptions and simple workarounds are recorded in this brief.')
h('Required capabilities')
bullet('Scan the Delivery Note and show its article numbers and expected quantities in a readable receiving list.')
bullet('Use the master lookup to show the product identity for a scanned or manually entered article number.')
bullet('Receive cans by scanning their supplied labels. Each accepted one-can scan adds one to received quantity and available locker stock.')
bullet('Issue cans by scanning the same labels in Issue mode. Each accepted one-can scan adds one to issued quantity and deducts one from available locker stock.')
bullet('Show the process and the inventory quantity before and after each movement. Keep the inventory result available when switching operations.')
h('Simple process flow')
code([
    'Scan DN -> Expected article list -> Receive can label -> Stock +1',
    'Locker inventory -> Issue same can label -> Stock -1',
    'Unclear label -> Manual article entry -> Same Receive or Issue action',
])
h('Meaningful data display')
table(['Article', 'Product from master', 'DN qty', 'Received', 'Issued', 'Available'], [
    ('Article A', 'Matched sample paint', '8', '0', '0', '0'),
    ('Article B', 'Matched sample paint', '4', '0', '0', '0'),
], [1.0, 2.2, .8, 1.0, .8, 1.2])
p('The figures above are illustrative; use the supplied DN, labels and agreed master entries in the demo. DN quantity is expected delivery data. It does not create stock simply because the paper was scanned.')
p('Available stock = opening stock + received cans - issued cans. The selected operation determines the sign of the movement; the label only identifies the article.')

page('1 Demo walkthrough')
step('Prepare the supplied DN, matching sample can labels, master lookup and a known opening balance. Use one locker room and one active delivery. Set opening stock to zero for the example below.')
step('Scan the Delivery Note. Display its delivery number and expected article quantities, enriched with product names from the master. Correct any DN OCR values manually before the demo continues.')
step('Select Receive cans. Keep the selected DN and operation visible. Point the camera at the article number on the supplied can label.')
step('Read the article, resolve the exact master entry and accept one can. Show the matched product and increase received quantity and available stock by one.')
step('Scan another supplied label in Receive mode and show the corresponding article row updating. Switching articles preserves the inventory totals.')
step('Open the locker inventory view. Show DN expected quantity, cumulative received quantity, cumulative issued quantity and current available stock.')
step('Select Issue cans. Scan the same label used for the first received can and resolve the same master article.')
step('Accept the issue of one can. Show issued quantity increasing by one and available stock decreasing by one. DN expected quantity and received history remain unchanged.')
step('Show the full process flow and the movement result. If a label is unclear, demonstrate manual article entry instead of an OCR recovery workflow.')
h('Visible inventory proof for Article A')
table(['Demo stage', 'DN qty', 'Received', 'Issued', 'Available'], [
    ('After DN scan only', '8', '0', '0', '0'),
    ('After receiving one can', '8', '1', '0', '1'),
    ('After issuing the same article', '8', '1', '1', '0'),
], [3.0, 1.0, 1.0, 1.0, 1.0])
p('The demo may show a partial receipt such as 1 of 8 received. That is progress, not an automatic shortage exception. The number of physical demo cans does not need to equal the entire DN quantity to demonstrate the movement flow.')
p('The same physical can can enter and then leave the locker in the demonstration. Identical labels on different cans legitimately identify the same article; the MVP does not attempt individual can serialization.')

page('2 Demo data assumptions and workarounds')
table(['Assumption', 'MVP basis', 'Simple workaround'], [
    ('Supplied samples', 'The demo uses the supplied DN and labels on the actual cans.', 'Prepare the samples and opening quantities before presenting.'),
    ('Article identity', 'The printed can article resolves to one master entry and matches the receiving DN article.', 'Correct the sample master or prepare an explicit code mapping before the demo. No fuzzy matching workflow.'),
    ('Known products', 'Every demonstrated article already exists in the master lookup.', 'Choose the correct existing master article manually. Add missing demo reference data during setup, not while scanning.'),
    ('DN row matching', 'The active demo DN has one unambiguous receiving target per article.', 'Prepare the displayed demo list with any identical repeated rows consolidated. Multi-line allocation is outside the MVP.'),
    ('Counting unit', 'One physical can equals one counted piece for the selected article.', 'Use compatible sample cans and DN units. Kit, component and litre conversions are outside the MVP.'),
    ('Readable article', 'The happy path has a readable article region in the supplied label.', 'If unclear, enter the article number manually. The MVP does not repair the label or provide special OCR recovery.'),
    ('Product description', 'The master supplies a meaningful product name and, if available, pack size.', 'Correct the demo master during setup; keep the article number visible even if optional description data is absent.'),
    ('Opening balance', 'The starting quantity for each demonstrated article is known. The walkthrough starts at zero.', 'Seed a stated opening balance or reset controlled demo data before another run.'),
], [1.3, 2.65, 3.05])
p('Article numbers are lookup identifiers and should preserve their printed format, including leading zeros. If a supplied can uses a manufacturer product code rather than the DN article, prepare that sample mapping in the master before the demonstration.')
p('Manual entry is a normal MVP input path. It resolves the same master record and uses the same Receive or Issue action as a successful scan; it does not introduce another counting process.')

page('3 Operating assumptions and workarounds')
table(['Assumption', 'MVP basis', 'Simple workaround'], [
    ('One locker and delivery', 'Use one locker room and one active DN for receiving.', 'Keep the demo in that context. Cross-locker transfers and choosing among several open DNs are outside scope.'),
    ('One operator', 'One operator performs the demonstrated movements.', 'Run the demo sequentially. Concurrent counting and synchronization conflicts are outside scope.'),
    ('One can per accepted scan', 'A deliberate accepted scan represents one can, not every OCR preview frame.', 'Repeat accepted scans for additional cans. Manual quantity entry can handle an agreed group where needed.'),
    ('Explicit movement direction', 'Receive adds; Issue deducts. The mode is chosen before scanning and is displayed.', 'Cancel or correct the article before accepting a movement if the wrong context or code was selected.'),
    ('Stock available for issue', 'The demo issues cans that have already been received or are in opening stock.', 'Use a valid issue quantity. A basic guard prevents stock going below zero; no shortage-resolution workflow is required.'),
    ('Controlled demo inventory', 'The demo uses a small persisted inventory store keyed by article. Live production ERP integration is not assumed.', 'Use prepared demo data and display the stored balance. Keep this separate from a claim that a production ERP was updated.'),
    ('Simple corrections', 'Exceptions are handled through manual input or demo data preparation.', 'Use the existing simple quantity controls where appropriate or reset the controlled demo data. Do not build an exception subsystem.'),
    ('Partial receiving is visible', 'Received quantity can remain below DN expected quantity while cans are still arriving.', 'Show expected and received quantities separately. Automatic shortage, excess and variance approval are outside this demo path.'),
], [1.3, 2.65, 3.05])
h('Outside this MVP')
p('Stocktake modes, individual can tracking, FIFO or expiry allocation, complex unit conversions, multi-location transfers, concurrent operators, automated duplicate-can detection, advanced OCR recovery and production inventory synchronization are outside the MVP.')
p('The demo needs the supplied DN, readable sample can labels, matching article master entries and agreed starting quantities. These are preparation inputs rather than additional application workflows.')

page('4 Minimum implementation and demo checks')
h('Keep the implementation small')
table(['Area', 'Minimum change'], [
    ('DN display', 'Reuse the existing DN scanner and source corrections. Show expected quantities and master product names in a clean list.'),
    ('Can article capture', 'Add a can label action that reads only the article needed for lookup. Reuse the on-device OCR engine; provide manual article entry.'),
    ('Master lookup', 'Use the prepared article-keyed sample master or its explicit sample mapping. The existing UPC-only portal lookup is not the article matcher.'),
    ('Inventory movements', 'Add Receive and Issue operations to one controlled locker inventory. Persist each accepted movement once and update the displayed balance.'),
    ('Data separation', 'Keep DN expected quantity, cumulative received quantity, cumulative issued quantity and available stock separate. Displaying or finalizing a DN must not add its quantity to stock again.'),
], [1.45, 5.55])
p('Stock issue is a new inventory movement. Decreasing the DN physical receiving counter is not the outgoing-can flow: it would change received history instead of recording a can leaving the locker. The new Issue operation must keep that history intact.')
p('Only article identification is needed from the can for this demo. Full paint-label metadata, UPC, batch number and expiry extraction should not become mandatory steps in the movement flow.')
h('Demo acceptance')
bullet('The supplied DN is scanned and its expected items are displayed meaningfully.')
bullet('A supplied can label is read, matched to the master and counted into stock in Receive mode.')
bullet('A second article can be received without losing the first article quantity.')
bullet('The same label is scanned in Issue mode and its available stock decreases by one.')
bullet('DN expected quantity and received history do not decrease when a can is issued.')
bullet('Manual article entry performs the same receipt or issue when the label is unclear.')
bullet('The stored before and after quantities are visible, and the flow is explained from DN to receipt to issue.')
p('This brief records the revised MVP requirements and the demonstration to build. The existing app provides DN scanning and manual receiving controls; can-driven article lookup and the stock issue flow are implementation additions.')

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)

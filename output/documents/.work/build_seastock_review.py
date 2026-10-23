from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path('/home/suchithgoud/Downloads/Label_Reader')
OUT = ROOT / 'output/documents/Seastock_Workflow_and_Can_Scanning_Review.docx'
doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(.65)
section.bottom_margin = Inches(.65)
section.left_margin = Inches(.75)
section.right_margin = Inches(.75)
section.footer_distance = Inches(.25)

for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3', 'List Number', 'List Bullet', 'List Paragraph']:
    style = doc.styles[name]
    style.font.name = 'Arial'
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.widow_control = True
for style in doc.styles:
    for border in list(style.element.xpath('.//w:pBdr')):
        border.getparent().remove(border)
normal = doc.styles['Normal']
normal.font.size = Pt(11)
normal.paragraph_format.line_spacing = 1.10
normal.paragraph_format.space_after = Pt(7)
doc.styles['Title'].font.size = Pt(25)
doc.styles['Title'].font.bold = True
doc.styles['Title'].paragraph_format.space_after = Pt(10)
doc.styles['Subtitle'].font.size = Pt(12)
doc.styles['Subtitle'].paragraph_format.space_after = Pt(8)
for name, size in [('Heading 1', 18), ('Heading 2', 12.5), ('Heading 3', 11.5)]:
    doc.styles[name].font.size = Pt(size)
    doc.styles[name].font.bold = True
    doc.styles[name].paragraph_format.space_before = Pt(10)
    doc.styles[name].paragraph_format.space_after = Pt(7)
    doc.styles[name].paragraph_format.keep_with_next = True

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
run = footer.add_run('Pixorex Seastock review   ')
run.font.size = Pt(9)
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)
doc.core_properties.title = 'Seastock workflow and physical can scanning review'
doc.core_properties.subject = 'Current Pixorex workflows and proposed article based field counting'
doc.core_properties.author = 'Pixorex'
doc.core_properties.keywords = 'Seastock, Pixorex, article master, physical stock, delivery note, paint cans'

numbering = doc.part.numbering_part.element
abstract_id = max(int(node.get(qn('w:abstractNumId'))) for node in numbering.findall(qn('w:abstractNum'))) + 1
abstract = OxmlElement('w:abstractNum')
abstract.set(qn('w:abstractNumId'), str(abstract_id))
level = OxmlElement('w:lvl')
level.set(qn('w:ilvl'), '0')
for tag, value in [('start', '1'), ('numFmt', 'decimal'), ('lvlText', '%1.'), ('lvlJc', 'left')]:
    child = OxmlElement('w:' + tag)
    child.set(qn('w:val'), value)
    level.append(child)
indent_props = OxmlElement('w:pPr')
indent = OxmlElement('w:ind')
indent.set(qn('w:left'), '300')
indent.set(qn('w:hanging'), '300')
indent_props.append(indent)
tabs = OxmlElement('w:tabs')
tab = OxmlElement('w:tab')
tab.set(qn('w:val'), 'num')
tab.set(qn('w:pos'), '300')
tabs.append(tab)
indent_props.insert(0, tabs)
level.append(indent_props)
abstract.append(level)
numbering.insert(list(numbering).index(numbering.find(qn('w:num'))), abstract)
current_num_id = None

def p(text, bold=False):
    para = doc.add_paragraph()
    para.add_run(text).bold = bold
    return para

def h(text, level=2):
    return doc.add_heading(text, level)

def step(text):
    global current_num_id
    if current_num_id is None:
        current_num_id = max(int(node.get(qn('w:numId'))) for node in numbering.findall(qn('w:num'))) + 1
        num = OxmlElement('w:num')
        num.set(qn('w:numId'), str(current_num_id))
        reference = OxmlElement('w:abstractNumId')
        reference.set(qn('w:val'), str(abstract_id))
        num.append(reference)
        override = OxmlElement('w:lvlOverride')
        override.set(qn('w:ilvl'), '0')
        start = OxmlElement('w:startOverride')
        start.set(qn('w:val'), '1')
        override.append(start)
        num.append(override)
        numbering.append(num)
    para = doc.add_paragraph(text, style='List Number')
    props = para._p.get_or_add_pPr().get_or_add_numPr()
    props.get_or_add_ilvl().val = 0
    props.get_or_add_numId().val = current_num_id
    para.paragraph_format.space_after = Pt(6)
    return para

def bullet(text):
    para = doc.add_paragraph(text, style='List Bullet')
    para.paragraph_format.space_after = Pt(5)
    return para

def page(title):
    global current_num_id
    current_num_id = None
    doc.add_page_break()
    h(title, 1)

def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for col, width in zip(t.columns, widths):
        col.width = Inches(width)
    for cell, label, width in zip(t.rows[0].cells, headers, widths):
        cell.width = Inches(width)
        cell.text = label
    for row in rows:
        cells = t.add_row().cells
        for cell, text, width in zip(cells, row, widths):
            cell.width = Inches(width)
            cell.text = text
    for index, row in enumerate(t.rows):
        row_properties = row._tr.get_or_add_trPr()
        if index == 0:
            repeat = OxmlElement('w:tblHeader')
            repeat.set(qn('w:val'), 'true')
            row_properties.append(repeat)
        cant_split = OxmlElement('w:cantSplit')
        row_properties.append(cant_split)
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props = cell._tc.get_or_add_tcPr()
            margins = OxmlElement('w:tcMar')
            for edge, value in [('top', 95), ('bottom', 95), ('left', 110), ('right', 110)]:
                node = OxmlElement('w:' + edge)
                node.set(qn('w:w'), str(value))
                node.set(qn('w:type'), 'dxa')
                margins.append(node)
            props.append(margins)
            borders = OxmlElement('w:tcBorders')
            for edge in ['top', 'left', 'bottom', 'right']:
                node = OxmlElement('w:' + edge)
                node.set(qn('w:val'), 'single')
                node.set(qn('w:sz'), '4')
                node.set(qn('w:color'), 'D9D9D9')
                borders.append(node)
            props.append(borders)
            shade = OxmlElement('w:shd')
            shade.set(qn('w:fill'), '0B1F44' if index == 0 else ('F1F5F9' if index % 2 == 0 else 'FFFFFF'))
            props.append(shade)
            for para in cell.paragraphs:
                para.paragraph_format.space_after = Pt(0)
                para.paragraph_format.line_spacing = 1.05
                for run in para.runs:
                    run.font.name = 'Arial'
                    run.font.size = Pt(10)
                    run.font.bold = index == 0
                    run.font.color.rgb = RGBColor.from_string('FFFFFF' if index == 0 else '000000')
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t

def code(lines):
    for line in lines:
        para = doc.add_paragraph()
        para.paragraph_format.space_after = Pt(1)
        para.paragraph_format.line_spacing = 1.0
        run = para.add_run(line)
        run.font.name = 'Liberation Mono'
        run.font.size = Pt(10)
    para.paragraph_format.space_after = Pt(9)

doc.add_paragraph('Seastock workflow and physical can scanning review', style='Title')
doc.add_paragraph('Pixorex project   2 October 2026   For operations and development review', style='Subtitle')
p('We recommend a field flow in which an operator scans the article number on a paint can, resolves it against the business article master, and records a deliberate physical quantity in a count session. A delivery receipt should carry a selected DN context. A stocktake of existing stock should work without a DN.')
p('The product master establishes what the product is. The DN establishes what a particular delivery should contain. The count session establishes how many physical units were observed, where, and for which task. An article number alone cannot identify the correct delivery.')
h('Current implementation and proposed extension')
table(['Capability', 'Current implementation', 'Field counting requirement'], [
    ('Can label extraction', 'Photos, cloud field extraction, operator review and approved records.', 'Focused article OCR and master identification.'),
    ('Inventory lookup', 'Approved paint records grouped by UPC and compared with a seeded inventory table.', 'Canonical article lookup and real location stock integration.'),
    ('DN receiving', 'On-device DN OCR, source review, saved physical counters and local final form.', 'Route accepted can scans into receiving counts.'),
    ('Stock posting', 'DN submission is local; the portal lookup does not post stock movements.', 'Reviewed, authorized and idempotent stock commit.'),
], [1.25, 2.9, 2.85])
h('Workflow boundaries')
code([
    'CURRENT  Can photos -> Review and cloud approval -> UPC lookup display',
    'CURRENT  DN pages   -> DN check and physical count -> Local final form',
    'PROPOSED Can article -> Master and scoped count -> Review and stock commit',
])
p('This review covers the implemented mobile and portal paths, the proposed field counting flow, DN reconciliation rules, stock update semantics, implementation gaps and decisions required before development. The source baseline is Pixorex 1.0.4 build 34 on the feature/ppg-delivery-note-scanner working tree. Source references are listed on the final page.')

page('1 Current paint label and portal flow')
step('The operator identifies themselves by email, opens the general camera action, and selects Paint / Chemical Label from the active schema catalog. Schemas and prompts are loaded from Supabase.')
step('The operator captures or imports up to ten views of the same product. Front, side and back views are combined into one logical extraction, including product code, pack size, batch information and UPC when visible.')
step('The mobile app sends those images to the configured cloud extraction service. It displays extracted values, confidence and review flags, missing mandatory fields, and any schema mismatch for operator review.')
step('The operator corrects or confirms the fields and submits. The app saves locally, uploads images and extraction data to Supabase, writes the approved values, and assigns the configured client for portal visibility.')
step('The client signs into the portal to view approved records, source images, search results and CSV exports. The Inventory Lookup view groups eligible paint records and compares them with the inventory reference table.')
h('How the existing lookup works')
p('paint_basket_view groups approved paint_label records by client and normalized UPC. It calculates scanned_qty as the number of approved records, then joins paint_inventory using UPC. There is no product_code or DN article fallback. Approved records with no valid UPC are absent from this lookup, even though they can remain visible in Records. [S1-S4]')
p('The inventory table contains seeded demonstration quantities and product attributes. The view reads available_qty and derives stock status; submitting a scan does not increase or replace available_qty. Historical approved scan records therefore cannot establish a scoped stock balance. The portal caption Paint cans counted onsite needs different semantics for production counting. [S4-S5]')
h('Runtime and alternate backend boundaries')
p('The general label path needs connectivity, active schema/prompt data and API configuration. Pixorex build 34 was built without a label-scan API key; its generic label extraction path requires configuration before use. Local saving and merged history do not constitute a durable inventory posting queue.')
p('The repository also contains a FastAPI path with extraction, correction, Excel persistence and optional Google Sheets synchronization. It is an alternative extraction/export path, not the active AppState cloud path and not a stock movement service. [S6]')

page('2 Current DN receiving and physical verification flow')
step('Start or resume a delivery-note session from the truck action. Capture or import all pages in any order. Accepted scans return to the page checklist so scanning can continue.')
step('Check the delivery identity, printed page numbers and declared page count. Review every source page against the paper, correct article numbers or printed quantities, resolve unreadable rows, and check full-page coverage.')
step('Review the printed Piece summary against the DN item total where available. Correct the summary or acknowledge a genuine source discrepancy. Quantity-less headings and batch details remain context, not additional receiving rows.')
step('Open the consolidated receiving list when the DN source checks are complete. Search articles and work in any arrival order. Printed DN quantities are expected values, not automatically verified physical counts.')
step('Verify each physical quantity. Running changes are saved locally. Switching articles needs no Pause action; the unfinished total remains available under Counting.')
step('Finish each count, record a reason for any final variance, review the final form, and submit locally. The submitted snapshot is read-only and is not sent to the portal or an inventory system. [S7-S9]')
table(['Operator action', 'Effect'], [
    ('Matches DN', 'Explicitly confirms the DN quantity after checking the stock.'),
    ('Start count and +1 or -1', 'Starts from zero and changes a running count one unit at a time. Minus stops at zero.'),
    ('Add batch', 'Adds the number just encountered to the existing total.'),
    ('Set total by tap or long press', 'Replaces the running total with the entered physical total.'),
    ('Finish count and Reopen count', 'Verifies the item, or resumes it later while retaining its total.'),
], [2.3, 4.7])
p('The stored states are To check, Counting, Verified and Recheck. Source page confirmation and physical confirmation are separate. A DN quantity correction preserves an entered physical total and requires rechecking. An article correction also requires reconciling which article the existing count belongs to.')
p('Repeated DN article occurrences remain separate source lines. Once a page has physical counts, direct DN corrections remain available but replacing or re-running OCR on that page is blocked to protect the counts. The current receiving controls are manual; there is no Scan can article action feeding them. [S8-S10]')

page('3 Proposed physical can scanning flow')
step('Create a count session for a client and vessel/site location. Choose Receiving or Stocktake. For Receiving, select the delivery context before quantities are attributed to a DN; allow unassigned observations only as a pending exception.')
step('Point the camera at the article-number area on one physical can. Read the printed article candidate. Keep the source crop and raw text available for confirmation or correction.')
step('Resolve the candidate against the authoritative article master or a permitted cached snapshot. Show the canonical article, product description, pack size and counting unit. A manufacturer product code or UPC may resolve through an approved alias mapping.')
step('Confirm the product and count one can, or explicitly add a group of identical cans. An accepted one-can event adds one unit; Add batch adds the entered number of cans. Extra views of the same can support that event and do not add units.')
step('Save the count event locally, then continue scanning another article. Resume any unfinished article from search or Counting without losing its running total. Provide Undo last count and direct total correction with an audit reason.')
step('If the session is Receiving, reconcile observed articles against the selected DN using the rules on the next page. Keep unresolved identities, unit mismatches and unassigned deliveries out of automatic stock posting.')
step('Review totals and exceptions, finish the session, and submit it to the stock service. Show Stock updated only after the server accepts the commit. A saved device draft is not yet a posted stock update.')
h('Capture guidance for cylindrical cans')
bullet('Bring the printed article region toward the center of the image. Rotate the can or move around it for the clearest view instead of trying to flatten the entire cylindrical label with a DN page crop.')
bullet('Avoid glare, the label seam, blur and steep angles. Retake a close view or enter the code manually when the article is unreadable.')
bullet('Keep article, product code, color code, batch number and UPC distinct. Preserve leading zeros and the master identifier format; ambiguous characters require confirmation.')
p('The existing on-device OCR engine can be reused to read a label region. The DN parser, page identity rules and page completeness checks should not be used as the can-label parser or as stocktake prerequisites. Field OCR reliability still requires a pilot using real cans, labels and conditions. [S11]')

page('4 Matching a physical can count to a DN')
p('Matching has two stages. First, the article master identifies the product. Second, the selected receiving session identifies the delivery. Matching by article alone cannot distinguish two open DNs, old stock from new arrivals, or different pack sizes for the same product family.')
table(['Situation', 'Proposed handling'], [
    ('One compatible article entry in the selected DN', 'Link the count to that entry after matching the canonical article, pack and unit.'),
    ('Repeated compatible article entries in one DN', 'Show a combined expected total only when aggregation is approved. Retain every source occurrence and apply the agreed reconciliation/allocation policy.'),
    ('The article appears in several open DNs', 'Use the DN selected for the session or require an explicit choice. Do not choose a delivery from the article code alone.'),
    ('Master match exists but the selected DN has no compatible article', 'Retain an unexpected-article exception for review. Do not silently add it to the DN or post it as an expected receipt.'),
    ('Article is unreadable, unknown or ambiguous', 'Retake, enter the code manually, or resolve a master candidate. The quantity remains unresolved until product identity is confirmed.'),
    ('No DN is selected', 'A Stocktake can proceed within its location scope. A Receiving observation stays unassigned until its delivery is identified.'),
], [2.35, 4.65])
h('Units and source line identity')
p('A physical can count must use a defined unit. DN PC may mean one can, a component or a set. A paint kit can contain several cans, and pack sizes can differ. Count cans first and convert to DN units only through an approved master rule. Do not infer kit conversion, paint volume remaining, or usable contents from the article number alone.')
p('The current DN model stores one physical count per printed occurrence. Grouping counts by canonical article, pack and unit would be a new behavior. For identical repeated rows, reviewers should approve aggregate reconciliation while retaining source provenance. Where batch, component or PO allocation matters, request an explicit allocation rather than distributing cans arbitrarily.')
p('Stock already mixed with earlier deliveries cannot be attributed to a new DN simply by reading its article. Receiving needs a segregated delivery area, known receipt batch or another agreed operational way to identify newly arrived stock.')

page('5 Count quantities and stock update semantics')
p('We should maintain product identity, count observations and stock balances separately. A count is an observation in a defined task and location. The inventory balance changes only through the accepted receipt or stocktake process.')
table(['Record', 'Minimum proposed content'], [
    ('Article master', 'Canonical article ID, approved printed-code aliases, description, variant/pack, unit, active status and master version.'),
    ('Count session', 'Session ID, task mode, client, site/vessel and location, count coverage, operator, status, start time and optional delivery/DN reference.'),
    ('Count event', 'Unique event ID, session and article IDs, quantity delta or total correction, unit, source images/raw OCR, operator and timestamp.'),
    ('DN source occurrence', 'Delivery ID, page/capture and occurrence IDs, verified printed article, expected quantity, unit and source evidence.'),
    ('Receiving reconciliation', 'Observed article count, compatible DN occurrence links, approved unit conversion, variance and exception/allocation decision.'),
    ('Stock commit', 'Unique commit ID, mode/scope, accepted quantities or adjustment, balance/version check, approval and server acknowledgment.'),
], [1.9, 5.1])
h('Illustrative quantity examples')
p('Receiving: a location has 100 cans. A DN expects 8 and the delivery contains 7. The receiving count is 7 and the receipt variance is -1. After accepting this receipt once, the location balance becomes 107; the count does not replace the balance with 7.')
p('Stocktake: the recorded location balance is 100, but a complete count finds 97. The observed stocktake total is 97 and the proposed adjustment is -3 after review. It is not an inbound receipt and does not add 97 to the existing balance.')
p('Mixed shelf: observing 107 cans with the same article does not prove that 7 came from a particular DN. That inference requires a controlled receiving context or traceable receipt evidence.')
p('These examples assume whole, compatible can units. Define whether empty, damaged, open or partially used containers belong in the count, and whether batch/expiry separation is required. Counting containers does not measure remaining paint volume.')

page('6 Proposed field screens and operator behavior')
h('Start a count session')
code([
    'Physical stock counting',
    'Task       [ Receiving ]  [ Stocktake ]',
    'Location   [ Select vessel or site and store ]',
    'Delivery   [ Select DN ]   Receiving only',
    '                              [ Start session ]',
])
p('Receiving binds the count to a delivery context. Stocktake covers existing stock at the selected location and has no mandatory DN step. Show the task and location throughout the session so counts from different purposes are not mixed.')
h('Scan and confirm a can')
code([
    'Scan can article        Receiving or Stocktake',
    '                 [ Camera view of article ]',
    'Article A   Product and pack from the master',
    'Physical total in this session    5 CAN',
    '[ Count this can +1 ]  [ Add batch ]  [ Retake ]',
])
p('Count this can adds one confirmed physical unit. Add batch is for several identical cans encountered together; entering 5 adds five to the current total. Retake or additional views belong to the same can-reading event until a unit is accepted.')
h('Consolidated running list')
code([
    'Search article ...      [ To check ] [ Counting ] [ Verified ]',
    'Article A                 5 CAN   Counting',
    '[ -1 ] [ 5 CAN  Edit ] [ +1 ] [ Scan more ]',
    '[ Add batch ]                            [ Finish count ]',
    'Article B                 3 CAN   Counting',
])
p('The operator may count A = 5, switch to B = 3, and later add 2 to A so A becomes 7. Each change remains saved within the same scoped session. There is no Pause action. Edit or long press sets the total; entering 5 replaces the total with 5. Reopen retains an already verified total if more stock is subsequently found.')
p('Seed receiving rows from the DN and stocktake To check rows from scoped book inventory where available. Unscanned articles stay unchecked; zero requires explicit physical verification. A partial count covers only its declared area/articles and cannot reset every other stock balance.')
p('Receiving adds a DN expected quantity and signed variance to these cards. Stocktake shows an observed total and, when available, a separate recorded balance. Final review shows all unresolved articles, allocations and reasons, then one clear submit action. DN source corrections remain in the paper review flow.')

page('7 Offline behavior and counting integrity')
h('Saving and synchronization')
p('Cache the authorized article master and, for receiving, the selected delivery data before field work. Record the master version and location scope. Known permitted cached articles can be counted offline; unknown or expired/changed master mappings require resolution before posting.')
p('Persist each accepted count operation and its audit record together. Use a stable count-event ID for upload and retry. The server must recognize the same event/commit ID and apply it once. Local states should distinguish Counting, Verified, Pending upload, Posted and Needs resolution.')
p('The current DN counter transaction and serialized write queue provide a useful local persistence pattern. They do not supply an article-master cache, an inventory synchronization queue or a remote stock commit. The cloud label submission path also does not provide those capabilities. [S1, S8-S9]')
h('Preventing duplicate and conflicting counts')
bullet('Many images of one can are one logical unit. Continuous OCR frames do not increment stock. Count only a deliberate accepted unit event or an explicit group quantity.')
bullet('An article code identifies a SKU, not an individual can. Identical cans share the same code, so deduplicating by article would suppress legitimate counts. A genuine serialized can identifier could strengthen duplicate detection if one exists.')
bullet('A new photograph of the same physical can can still be counted twice. The operation needs a physical separation/marking rule or another agreed control; an event UUID alone cannot detect this real-world duplicate.')
bullet('Undo and set-total corrections preserve history and invalidate final verification. Replays after an upload timeout must not create another stock movement.')
bullet('If multiple operators count the same stock, assign non-overlapping locations or receipt portions, or review how concurrent totals are combined. Recheck stale stocktake balances before committing an adjustment.')
h('Confirmation of a stock update')
p('For receiving, post an accepted receipt against the selected delivery, linked article and compatible unit. For stocktake, post an approved scoped observation/adjustment through the authoritative inventory service. Validate identity, units, scope, permissions and balance conflicts at commit time. Keep the device draft until acceptance is confirmed. Submitted local forms and failed uploads must not be labeled Stock updated.')
p('Verified drafts may be reopened before posting. Posted sessions stay immutable. Later arrivals require a new receipt or reviewed adjustment for the additional units; do not post the earlier full running total again.')

page('8 Implementation changes required')
table(['Area', 'Current capability', 'Proposed work'], [
    ('Article identification', 'General paint metadata extraction; on-device OCR available in the DN module.', 'Dedicated can article reader, source crop, ambiguity handling and manual entry.'),
    ('Master resolution', 'Demo portal lookup joins approved records to UPC-keyed inventory.', 'Business article master adapter, approved aliases, variant/unit checks and permitted offline cache.'),
    ('Count scope', 'DN session contains delivery identity and per-occurrence quantities.', 'Separate stock count sessions with receiving/stocktake mode and location. Avoid DN validation as a stocktake prerequisite.'),
    ('Quantity controls', 'Saved manual increments, batch additions, total editing, finish and reopen.', 'Bind accepted can events to the count target; add scan-event grouping and Undo.'),
    ('DN reconciliation', 'Printed occurrences preserved; final validation is per line.', 'Selected DN linking, duplicate-line policy, unit conversion and unexpected-article resolution.'),
    ('Inventory posting', 'Local immutable DN form; cloud label approval and read-only lookup.', 'Authorized atomic stock commit, idempotent retry, conflict resolution and posting acknowledgment.'),
    ('Portal reporting', 'Historical approved UPC scan count and seeded availability.', 'Views for scoped observed counts, receipt totals, stocktake results, posting state and real balances.'),
], [1.35, 2.55, 3.1])
p('The first field slice should establish accurate article identification and saved session counts. DN matching and live balance posting should follow only with approved master, unit and inventory contracts. This sequencing can support a useful field pilot without presenting local counts as posted stock.')
p('Preserve existing DN submissions and extraction history. Introduce new field count/session records rather than retroactively treating every approved paint extraction as a unique physical can or overwriting old receiving snapshots. Reuse the existing counter behavior while separating can events from document source occurrences.')

page('9 Decisions and pilot acceptance')
h('Decisions for the reviewers')
table(['Decision', 'Agreement required'], [
    ('Operating purpose', 'Receiving, stocktake or both as separate modes. Define whether receiving can begin with a pending DN assignment.'),
    ('Master authority', 'Identify the article master owner, access/export/API, canonical identifier and mappings from can codes/UPCs to DN articles.'),
    ('Scope and units', 'Agree client, vessel/site, location, can/PC/set conversions, pack variants and treatment of damaged/open/empty cans.'),
    ('Counting trigger', 'Choose deliberate scan plus confirmation or an approved automatic +1 rule. Agree multi-view grouping, batch entry and real-world duplicate prevention.'),
    ('DN allocation', 'Agree handling of repeated article lines, several open DNs, component/batch restrictions and unexpected products.'),
    ('Posting authority', 'Identify the real inventory system, who can approve/post, stocktake adjustment semantics, offline policy and concurrent-count ownership.'),
], [1.65, 5.35])
h('Pilot acceptance checks')
bullet('Use real cylindrical cans across label positions, glare, seams, blur and rotations. Validate the article against the master and a human-confirmed reference. Set an acceptable error/confirmation rule before field rollout.')
bullet('Verify leading zeros, similar characters, unknown codes, inactive articles and several master candidates. No unresolved product should change posted stock.')
bullet('Confirm that three views of one can add one unit; five confirmed identical cans entered as a batch add five; repeated upload of an event adds nothing extra.')
bullet('Count A, switch to B, restart the app, then resume A. Check +1/-1, explicit zero, set-total replacement, Undo, Finish and Reopen. Verify that unscanned articles remain unchecked.')
bullet('Test one DN line, repeated DN lines, several open DNs, pack/set mismatches, unexpected articles and a stocktake with no DN.')
bullet('Lose connectivity during capture and during submission. Verify recovery, master-version checks, duplicate-safe posting and visible pending/posted states.')
p('The existing source checks, counter behavior, correction handling and submission locking have automated coverage; the most recent full project run passed 61 tests. That result does not validate cylindrical can OCR, article-master matching or live inventory posting, which are new requirements.')

page('10 Source references and implementation evidence')
p('The references below identify the current code behavior used for this workflow review. Proposed behavior in Sections 3 to 9 requires development and business agreement; it is not a capability of the current APK.')
sources = [
    ('S1', 'frontend/lib/services/app_state.dart', 'loadSchemas, extractFields and submitExtraction define the active generic label extraction, local save and Supabase approval path.'),
    ('S2', 'frontend/lib/services/openai_service.dart and frontend/lib/services/config.dart', 'extractFields performs the configured cloud vision call and checks API configuration; schema/confidence rules are separate from physical counting.'),
    ('S3', 'frontend/lib/screens/capture_screen.dart and results_screen.dart', 'The general camera groups up to ten images. Review permits explicit missing-field override and validates non-empty UPC format.'),
    ('S4', 'client_portal_poc/sql/01_complete_setup.sql Section 7', 'paint_inventory is seeded demo reference data. paint_basket_view filters valid UPCs, groups records, counts them and joins on UPC.'),
    ('S5', 'client_portal_poc/app.js loadPaintBasket and renderBasket', 'The portal reads scan totals and inventory availability; it does not post a quantity change from an article scan.'),
    ('S6', 'backend/main.py, excel_store.py and google_sheets.py', 'Alternative extraction, correction and export services; no implemented stock count or receipt posting endpoint.'),
    ('S7', 'frontend/lib/features/ppg_delivery_note/screens/ppg_session_screen.dart and ppg_page_review_screen.dart', 'Scan progression, page checklist, DN identity/content correction and coverage confirmation.'),
    ('S8', 'frontend/lib/features/ppg_delivery_note/services/ppg_session_controller.dart', 'Physical count state transitions, serialized count writes, source correction preservation and source/physical confirmation separation.'),
    ('S9', 'frontend/lib/features/ppg_delivery_note/persistence/ppg_repository.dart and domain/local_submission.dart', 'SQLite version 9, atomic count/audit commits, local immutable submission and readiness checks.'),
    ('S10', 'frontend/lib/features/ppg_delivery_note/parser/consolidator.dart and widgets/ppg_physical_quantity_control.dart', 'Printed occurrences remain distinct; quantity buttons, batch entry, total edit and variance reason dialogs.'),
    ('S11', 'frontend/lib/features/ppg_delivery_note/ocr/ppg_ocr_adapter.dart and android/app/src/main/kotlin/com/amzur/equiptrack/ppg/PpgOcrPlugin.kt', 'Existing native on-device Latin text recognition and region reading can supply the engine for a new article reader.'),
    ('S12', 'frontend/test/ppg_delivery_note/ppg_receiving_workflow_test.dart and other frontend tests', 'Current automated coverage includes persistence, rapid edits, zero, reasons, DN corrections, duplicate occurrences and submission locking.'),
]
for label, path, evidence in sources:
    para = p(label + '   ' + path, bold=True)
    for run in para.runs:
        run.font.size = Pt(9.5)
    para.paragraph_format.keep_with_next = True
    para = p(evidence)
    for run in para.runs:
        run.font.size = Pt(10)
    para.paragraph_format.space_after = Pt(7)

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
